from io import BytesIO
from html import escape

from django.conf import settings
from django.core.exceptions import ImproperlyConfigured
from django.db.models import Count, Q, Sum, Value
from django.db.models.functions import Coalesce
from django.http import FileResponse
from django.shortcuts import get_object_or_404
from django.views import View
from django.contrib.auth.mixins import LoginRequiredMixin

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from workers.models import Worker
from .models import Tabel

HEADERS = [
    "№",
    "Табельний номер",
    "ПІБ працівника",
    "Посада",
    "Відпрацьовані дні",
    "Пропуски",
    "Відпустка",
    "Лікарняний",
    "Відрядження",
]

def get_timesheet_data(tabel):
    """
    Формує дані для PDF та Excel.
    Використовує ту саму логіку, що й
    AccountantTimesheetDetailView.
    """

    workers = (
        Worker.objects.filter(unit=tabel.unit)
        .annotate(
            total_skip_days=Coalesce(
                Sum(
                    "tabel_records__skip_days",
                    filter=Q(tabel_records__tabel=tabel),
                ),
                Value(0),
            ),
            total_vacation=Coalesce(
                Sum(
                    "tabel_records__vacation",
                    filter=Q(tabel_records__tabel=tabel),
                ),
                Value(0),
            ),
            total_hospital=Coalesce(
                Sum(
                    "tabel_records__hospital",
                    filter=Q(tabel_records__tabel=tabel),
                ),
                Value(0),
            ),
            total_assigment=Coalesce(
                Sum(
                    "tabel_records__assignment",
                    filter=Q(tabel_records__tabel=tabel),
                ),
                Value(0),
            ),
            total_work_days=Coalesce(
                Sum(
                    "tabel_records__count_work_day",
                    filter=Q(tabel_records__tabel=tabel),
                ),
                Value(0),
            ),
            record_count=Count(
                "tabel_records", 
                filter=Q(tabel_records__tabel=tabel),
            )
        ).order_by("surname", "firstname", "lastname")
    )

    rows = []

    for number, worker in enumerate(workers, start=1):
        if worker.record_count:
            days = [
                worker.total_skip_days,
                worker.total_vacation,
                worker.total_hospital,
                worker.total_assigment,
            ]
        else:
            days = ["-"] * 4


        rows.append([
            number,
            worker.tabel_num,
            (
                f"{worker.surname} "
                f"{worker.firstname} "
                f"{worker.lastname}" 
            ),
            worker.post,
            worker.total_work_days if worker.record_count else "-",
            *days,
        ])

    return rows


class AccountantExportMixin(LoginRequiredMixin):
    """
    Доступ лише до табелів підрозділів,
    призначених поточному бухгалтеру.
    """

    def get_tabel(self):
        return get_object_or_404(
            Tabel.objects.select_related("unit"),
            pk=self.kwargs["pk"],
            unit__accountants=self.request.user,
        )


class AccountantTimesheetPDFView(AccountantExportMixin,View):
    def get(self, request, pk):
        tabel = self.get_tabel()
        rows = get_timesheet_data(tabel)

        buffer = BytesIO()

        font_path = (
            settings.BASE_DIR
            / "fonts"
            / "DejaVuSans.ttf"
        )

        if not font_path.is_file():
            raise ImproperlyConfigured(
                "Додайте DejaVuSans.ttf "
                "у папку fonts/ проєкту."
            )
        
        if "TimesheetFont" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(
                TTFont("TimesheetFont", str(font_path))
            )

        pdf = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            rightMargin=25,
            leftMargin=25,
            topMargin=30,
            bottomMargin=30,
        )

        normal = ParagraphStyle(
            "TimesheetNormal",
            fontName="TimesheetFont",
            fontSize=8,
            leading=12,
        )

        heading = ParagraphStyle(
            "TimesheetHeading",
            parent=normal,
            fontSize=15,
            leading=20,
            spaceAfter=10,
        )

        story = [
            Paragraph(
                f"Табель обліку робочого часу",
                heading,
            ),
            Paragraph(
                escape(
                    f"Підрозділ: {tabel.unit.title}"
                ),
                normal,
            ),
            Paragraph(
                f"Період: "
                f"{tabel.t_month:02d}.{tabel.t_year}",
                normal,
            ),
            Spacer(1, 18),
        ]

        data = [
            [
                Paragraph(escape(str(cell)), normal)
                for cell in HEADERS
            ]
        ]

        for row in rows:
            data.append([
                Paragraph(escape(str(cell)), normal)
                for cell in row
            ])

        table = Table(
            data,
            colWidths=[
                28, 75, 155, 105,
                70, 70, 70, 70, 70,
            ],
            repeatRows=1,
            hAlign="LEFT",
        )

        table.setStyle(TableStyle([
            (
                "BACKGROUND",
                (0, 0), (-1, 0),
                colors.HexColor("#E2E8F0"),
            ),
            (
                "VALIGN",
                (0, 0), (-1, -1),
                "MIDDLE",
            ),
            (
                "GRID",
                (0, 0), (-1, -1),
                0.4,
                colors.HexColor("#CBD5E1"),
            ),
            (
                "ROWBACKGROUNDS",
                (0, 1), (-1, -1),
                [colors.white, colors.HexColor("#F8FAFC")],
            ),
            (
                "LEFTPADDING",
                (0, 0), (-1, -1),
                6,
            ),
            (
                "RIGHTPADDING",
                (0, 0), (-1, -1),
                6,
            ),
            (
                "TOPPADDING",
                (0, 0), (-1, -1),
                8,
            ),
            (
                "BOTTOMPADDING",
                (0, 0), (-1, -1),
                8,
            ),
        ]))

        story.append(table)

        pdf.build(story)

        buffer.seek(0)

        filename = (
            f"timesheet_{tabel.unit.code}_"
            f"{tabel.t_month:02d}_{tabel.t_year}.pdf"
        )

        return FileResponse(
            buffer,
            as_attachment=True,
            filename=filename,
            content_type="application/pdf",
        )


class AccountantTimesheetExcelView(
    AccountantExportMixin,
    View,
):
    def get(self, request, pk):
        tabel = self.get_tabel()
        rows = get_timesheet_data(tabel)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Табель"


        sheet.merge_cells("A1:I1")
        sheet["A1"] = "Табель обліку робочого часу"

        sheet["A1"].font = Font(
            size=16,
            bold=True,
            color="FFFFFF",
        )

        sheet["A1"].fill = PatternFill(
            "solid",
            fgColor="1E3A5F",
        )

        sheet["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        sheet.row_dimensions[1].height = 35

        sheet["A2"] = "Підрозділ:"
        sheet["B2"] = tabel.unit.title

        sheet["A3"] = "Період:"
        sheet["B3"] = (
            f"{tabel.t_month:02d}.{tabel.t_year}"
        )

        for col, title in enumerate(HEADERS, start=1):
            cell = sheet.cell(
                row=5,
                column=col,
                value=title,
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                "solid",
                fgColor="2563EB",
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        sheet.row_dimensions[5].height = 32

        for row_index, row_data in enumerate(
            rows,
            start=6,
        ):
            for col_index, value in enumerate(
                row_data,
                start=1,
            ):
                cell = sheet.cell(
                    row=row_index,
                    column=col_index,
                    value=value,
                )

                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )

                if row_index % 2 == 0:
                    cell.fill = PatternFill(
                        "solid",
                        fgColor="F1F5F9",
                    )

            sheet.row_dimensions[row_index].height = 25

        widths = [8, 22, 35, 25, 22, 15, 15, 15, 17]

        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[
                get_column_letter(col)
            ].width = width

        last_row = max(5, 5 + len(rows))

        sheet.auto_filter.ref = f"A5:I{last_row}"

        sheet.freeze_panes = "A6"

        sheet.page_setup.orientation = "landscape"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.print_title_rows = "1:5"

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        filename = (
            f"timesheet_{tabel.unit.code}_"
            f"{tabel.t_month:02d}_{tabel.t_year}.xlsx"
        )

        return FileResponse(
            buffer,
            as_attachment=True,
            filename=filename,
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )
